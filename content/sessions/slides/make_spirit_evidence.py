#!/usr/bin/env python3
"""Spirituality in Sexuality - research brief as slides. One idea per slide.
SEFW palette, Belleza headings + Lora body, upright text."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
import os

NAVY =RGBColor(0x1D,0x28,0x3C); PLUM =RGBColor(0x74,0x23,0x4F)
BLUSH=RGBColor(0xE4,0xBB,0xC2); PINK =RGBColor(0xFC,0xE8,0xEA)
SAGE =RGBColor(0x7C,0xA1,0xA5); TEAL =RGBColor(0x22,0x46,0x52)
GOLD =RGBColor(0xC9,0xA8,0x6D)
HEAD_FONT="Belleza"; BODY_FONT="Lora"
SCHEME={"pink":(PINK,NAVY,PLUM),"blush":(BLUSH,NAVY,PLUM),"plum":(PLUM,PINK,GOLD),
        "navy":(NAVY,PINK,GOLD),"teal":(TEAL,PINK,GOLD),"sage":(SAGE,NAVY,PLUM),
        "gold":(GOLD,NAVY,PLUM)}

def new_deck():
    p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5); return p

def add(prs,bg,title,body=None,eyebrow=None,big=False,footer=True):
    bg_col,txt_col,accent=SCHEME[bg]
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=bg_col
    left=Inches(0.9); width=Inches(11.5); top=Inches(1.2)
    if eyebrow:
        eb=s.shapes.add_textbox(left,Inches(0.65),width,Inches(0.5))
        r=eb.text_frame.paragraphs[0].add_run(); r.text=eyebrow.upper()
        r.font.name=BODY_FONT; r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=accent
    tb=s.shapes.add_textbox(left,top,width,Inches(2.6))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    r=tf.paragraphs[0].add_run(); r.text=title
    r.font.name=HEAD_FONT; r.font.size=Pt(52 if big else 40); r.font.color.rgb=txt_col
    ln=s.shapes.add_shape(1,left,top+Inches(2.9 if big else 1.9),Inches(1.6),Pt(3))
    ln.fill.solid(); ln.fill.fore_color.rgb=accent; ln.line.fill.background(); ln.shadow.inherit=False
    if body:
        by=s.shapes.add_textbox(left,top+Inches(3.25 if big else 2.25),width,Inches(3.2))
        bf=by.text_frame; bf.word_wrap=True
        items=body if isinstance(body,list) else [body]
        for i,item in enumerate(items):
            para=bf.paragraphs[0] if i==0 else bf.add_paragraph()
            bullet=isinstance(body,list) and len(body)>1
            run=para.add_run(); run.text=("•  "+item) if bullet else item
            run.font.name=BODY_FONT; run.font.size=Pt(24 if bullet else 27)
            run.font.italic=False; run.font.color.rgb=txt_col
            para.space_after=Pt(12)
    if footer:
        ft=s.shapes.add_textbox(left,Inches(6.9),width,Inches(0.4))
        fr=ft.text_frame.paragraphs[0].add_run(); fr.text="SEXUALEMPOWERMENTFORWOMEN.COM"
        fr.font.name=BODY_FONT; fr.font.size=Pt(11); fr.font.color.rgb=accent
    return s

d=new_deck()

add(d,"navy","Spirituality in Sexuality",eyebrow="Radiant Women's Circle",
    body="What it means, how they connect, and the evidence.",big=True)

# What it means - one slide
add(d,"pink","What sacred sexuality means",eyebrow="What it means",
    body=["Bringing intention, presence and reverence to sex, so it's more than a physical act.",
          "A doorway to something larger than you: the divine, life force, oneness, love.",
          "The act doesn't change. The awareness you bring does. Communion, not performance."])

# How they connect - two slides
add(d,"teal","The same longing",eyebrow="How they connect",
    body=["Both spring from the same root: desire and longing.",
          "Sexuality: the longing to know and be known, to merge with another.",
          "Spirituality: the same longing, pointed at the whole, to belong to something greater.",
          "Two directions, one yearning, to end separateness."])
add(d,"plum","The body is the doorway",
    body=["Tantra never split them. Body-shameful, spirit-holy is the wound, not the truth.",
          "Transcendence, like pleasure, moves through the body, not around it.",
          "You reach the sacred by going in, not rising above. Bypassing can't get you there."])

# The evidence
add(d,"sage","Presence is the skill",eyebrow="The evidence · it's real, not woo",
    body="Mindfulness measurably improves women's desire, arousal and satisfaction (Dr Lori Brotto). The spiritual muscle and the sexual one are the same: presence.")
add(d,"navy","Union is literal",
    body="In eye-gaze studies, two people's brains, even their blinking, begin to synchronise, and they feel closer. Becoming one isn't only poetry.")
add(d,"gold","A chemistry of the sacred",
    body="Touch and intimacy release oxytocin, dopamine and endorphins: bonding, trust, dissolved boundaries, the afterglow.")
add(d,"blush","Vulnerability is the gateway",
    body="Being truly seen is tied to the deepest connection, the same vulnerability the traditions call surrender.")

add(d,"navy","The body isn't the obstacle",
    body="Received as a gift, it's the doorway to the sacred.",big=True)

os.makedirs("content/sessions/slides",exist_ok=True)
out="content/sessions/slides/Spirituality-in-Sexuality-Research.pptx"
d.save(out)
print("saved",out,len(d.slides._sldIdLst),"slides")
