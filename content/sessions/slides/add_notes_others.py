#!/usr/bin/env python3
"""Attach full facilitator-script presenter notes to the Emotions and
Needs/Desires decks, in slide order."""
from pptx import Presentation

EMO=[
"Welcome. Light the candle, read the circle agreements: what's shared here stays here; we witness, we don't fix; you can always pass; there's no such thing as too much. Soften them in: feet on the floor, soften jaw, shoulders, belly, three slow breaths. 'Tonight we learn to be with our feelings, safely, not running and not drowning.' One-word go-round: how are you arriving?",
"The whole frame for both sessions: lack of safety is the problem, not the emotions. Emotions aren't dangerous, being flooded is. We get flooded when we try to feel big things with no safety underneath.",
"So we start with safety. We only expand, feel more, open more, when we're not overwhelmed and not sliding into our old patterns. The moment overwhelm shows up, or you reach for the numbing thing, that's not failure, it's information: come back to safety first.",
"Give your definition of safety, slowly: I'm connected to myself; I feel what's good for me; I test it out, one small step; I read how the other responds; then I choose, step forward or step away. Moving like that, connected, checking, choosing, there's no overwhelm.",
"Land it: when you stay with yourself the whole way, there's no overwhelm, because you never abandon yourself to cope.",
"Pose the reflection questions before the breakout: When do I feel safe? When do I not? Where do I abandon myself to keep the peace? What does it feel like in my body when I leave myself?",
"Broadcast to the rooms. In pairs, Partner A about 5-6 minutes: 'One place I feel safe is... one place I don't is...' then 'One way I abandon myself is...' and 'What it feels like in my body when I leave myself is...' Partner B witnesses, no fixing. Swap. Main-room: one way I abandon myself.",
"Teach the numbing loop: stress makes emotions, emotions make stress; we reach for something outside ourselves, food, wine, phone, busyness, other people's problems, to fill an emotional hole.",
"And it can't be filled from outside. There's no signal that says your emotional need is met, no leptin for feelings. So the fix backfires, we feel guilty, and the feeling still waits to be heard.",
"Say it clearly: emotions are the messengers. Ignore them and they don't go away, they get louder, they start screaming until we listen. Numbing is what we do when we don't feel safe enough to feel, it's not a flaw.",
"One minute, soft music, journal: what do you reach for when a feeling gets too big? No judgement, just notice. Then: 'okay, completing this reflection.'",
"Dr Linda Bacon's reframe: if you numb, you don't have a problem with food or wine or your phone, you have a problem with taking care of yourself, and with not feeling safe enough to feel. The numbing falls away as you meet your needs and build safety. Are you with me?",
"Guide the practice, hands on heart: 'How do you feel? What's going on for you?' Listen. Then 'I hear you feel...' and name it. If it's hard: 'I see this is difficult. I am here.' Nothing to change, just being with yourself. One of the fastest ways back to safety.",
"Broadcast. In pairs, Partner A hands on heart, out loud to B: 'How do you feel? What's going on for you?' Answer honestly. Partner B's only words back: 'I hear you feel...' No advice. Swap. Main-room: what happened in your body when someone just listened?",
"Take-home: when a feeling rises or you catch yourself reaching to numb, first ask 'Am I safe? Am I connected to myself?' Then hands on heart, 'How do you feel?' and listen. That's the whole practice. Session 2 is what to do when the feeling is a wave.",
"Session 2. Soften them in and anchor safety, then recap: lack of safety is the problem, not the emotions. Go-round: one feeling I've been carrying this week. Pose the reflection questions for later: what feels too big? where am I braced? what brings me back to safety?",
"Before we go near a big feeling tonight, anchor safety: feet on the floor, hand on your heart, 'Right now, in this moment, I am here, and I am safe.'",
"Pose and let them sit: What feeling, if I let it, feels too big? Where do I feel it in my body, where am I braced? What brings me back to safety?",
"Not shutting it out, not drowning. Feeling into it without being taken over. The third way, right in the middle.",
"Grief comes in waves. You don't feel all of it at once; it rises, crests, passes for now, then returns. That's grief doing what grief does, not you failing.",
"The skill of self-regulation, titrating: go into the feeling for a breath or two, really feel it, then deliberately come back to safety, feet, breath, hand on heart. A little in, a little out. The opposite of numbing, and of drowning.",
"While the wave moves, keep part of your attention here: name five things you see, 'right now I am here and I am safe.' That's how you feel it without disappearing into it.",
"Safety first: tell them to pick something that's a 3 or 4 out of 10, not their biggest thing. Broadcast. In pairs, A brings a manageable feeling, feels it one or two breaths, says where it is; B guides A back to safety; a little in, a little out, three times; swap. B just holds the safety anchor. Main-room: what did you notice about coming back on purpose?",
"You can make peace in your head, and your body can still be holding it. The body lets go through breath, tears, moving, being held. Lead one minute: hand on the braced place, three long breaths, an audible sigh on each out-breath. Then soothe: hands on heart, 'I see this is hard. I am here.'",
"Introduce EFT, tapping, an easy and direct way to settle a big feeling and come back to safety in the body. Note you're a certified EFT practitioner and this is a quick version they can use right away.",
"Simply: tapping calms the amygdala's fight-or-flight alarm, shown on brain scans, so you shift into your logical brain and choose your response instead of being run by the feeling. It also helps the brain update: that's not a threat now, I'm safe.",
"Reassure about tapping on the 'negative': we don't attract more by naming it. We take the dirty laundry out and wash it instead of hiding it. Name it, clear it, and reframes come naturally.",
"Walk the steps: 1) choose a specific issue, rate it 1 to 10; 2) setup on the karate chop point, 'Even though I feel this, I love and accept myself,' three times; 3) tap the points saying your reminder phrase out loud; 4) breathe, re-rate, repeat until it drops. Sip water throughout.",
"Show the points on camera: top of head, eyebrow, side of the eye, under the eye, under the nose, chin, collarbone, under the arm at nipple level, wrist. If another issue surfaces mid-round, that's good, that's how we go deeper.",
"Broadcast, or lead together first. In pairs, each picks an issue, rates 1 to 10, says it; tap through the points together saying your own reminder phrase out loud; partner keeps time and company; breathe, re-rate; second round if still high; swap. If something big surfaces, invite them back to safety and offer one-to-one support, this is everyday self-help, not trauma treatment.",
"Take-home, three moves: 1) safety first, am I connected to myself; 2) hands on heart, how do you feel, listen; 3) tap when it's big, and with the waves, feel a little, come back, feel a little more.",
"Close: you don't master your emotions by controlling them, you master them by feeling safe enough to finally listen. Closing round, one line each: one feeling I'm going to stop numbing. Circle answers 'We hear you.' Blow out the candle.",
]

S1=[
"Welcome. Candle, circle agreements: what's shared here stays here; we witness, we don't fix or rescue; there's no right or wrong way; there's no such thing as too much. Soften in: feet, jaw, shoulders, belly, three breaths. Tonight we find her, your needs, wants, deepest desires. One-word go-round.",
"Frame the night: your needs, your wants, your deepest desires. We start by coming home to yourself.",
"Read the agreements aloud: what's shared here stays here; we witness, we don't fix or rescue; there's no right or wrong way; there's no such thing as too much in this room.",
"Guide: eyes closed, soften the jaw, drop the shoulders, unclench the belly, release the pelvic floor. 'She doesn't live in the tension, she lives in the softening.' Three slow breaths, let the day fall away.",
"Teach the distinction: Needs, what you can't be well without. Wants, the specific things that meet a need. Desires, the deep, alive pull, your compass. You can't ask for what you can't name, so we name them first.",
"You've been gathering everyone else's needs so long you lost the thread of your own. Tonight we pick it back up.",
"Claire Zammit's work. You are not overlooked by accident. Set it up gently, this is the piece most women have never heard.",
"The I'm Invisible pattern: my attention is on everyone but me; I assume you can see what I never said; I serve until I'm empty; I wait, quietly, to be discovered. The key mechanism: assuming it's others' job to see what we never make visible.",
"The self-fulfilling loop: believe you're invisible, disappear yourself while focused on others, they don't see you, you feel invisible, resentment, you ask with accusation, they defend, you feel invisible again. The belief builds its own proof.",
"Said gently: you have taught the people around you exactly how much of you to see. A habit, not a fact. Tonight we start practising the opposite.",
"The shift, presencing: I stop waiting to be seen, I turn toward myself and make myself visible. It is my destiny to be visible. From making others responsible for seeing me, to seeing and presencing myself.",
"Broadcast. In pairs: 'One way I make myself invisible is...' / 'One way I've trained people not to see my needs is...' / 'The story underneath might be... invisible, not enough, too much, a burden.' Partner witnesses, no fixing. Swap. Main-room: the way I disappear is...",
"Part one. The no is the doorway into your desires. Your body already knows it.",
"Journal round, about 7 minutes, soft music: What I'm tired of pretending is... / What I'm done tolerating is... / What I no longer want is... Feel it, don't think it.",
"Part two. Every no is holding a yes behind it.",
"Journal, then say it out loud: The kind of touch I'm craving is... / A way I want to receive more is... / If I trusted myself completely, I would... Then speak one want out loud in dyads.",
"Part three. Explore in pairs. Underneath the want is the desire your whole life is organised around. Guide a short drop-in first if there's time.",
"And if I had that, what would it give me? Keep asking, going seven levels deep. Reach for the true answer, not the impressive one.",
"Claiming round: 'I'm a radiant woman, and I desire...' The circle answers: You're allowed.",
"The power statement: I see myself. I am present to my own feelings, needs and desires. It is my destiny to be visible, and I take my rightful place.",
"Affirmation, hand on heart, together: I'm powerful. I can create life. I can create the life I desire. The power is in my hands.",
"Take-home: say one clean want a day, out loud; each night, whisper your deepest desire to yourself; you're teaching your body that your desire is safe. Closing round.",
]

S2=[
"Session 2. Soften in, recap Session 1: you found your needs; tonight you give them a voice, not a complaint, not a hint, a desire spoken so it can be met. Go-round.",
"Tonight we give her a voice. Frame: asking in a way that inspires.",
"Go-round: one desire I found last time was...",
"Alison Armstrong: you live in diffuse awareness, aware of everything at once; he lives more in single focus, one thing at a time. So when you share, he hears information, not a request. Neither of you is broken.",
"A hundred diffuse hints land as noise. One clear want lands as a request.",
"A real request has three things: specific (he can picture it); a by-when (it lands in time); room for a yes or a no (a request, not a demand). Nagging is repeating a vague thing; a real request is clean.",
"Broadcast. In pairs: share it the messy way first; then say it in one clean sentence, 'I want... by...'; partner reflects back the clear request. Swap. Main-room: what did it feel like to say the clean version?",
"Armstrong's four steps: 1) describe exactly what you want; 2) say what it would give you; 3) say what's in it for him; 4) ask 'what do you need from me to give me that?' That last line turns a request into a team.",
"Frustration is the gap between an expectation you never spoke and the influence you didn't use. An unspoken expectation is a resentment in waiting; a spoken request is an act of intimacy.",
"Broadcast. In pairs, build your ask in all four steps; partner receives it warmly in character and answers step 4. No apologising, say it like it's allowed. Swap. Main-room: what surprised you about asking that way?",
"A demand makes him wrong and he stops trying; an invitation hands him a way to be your hero. Assume he wants to give; ask for the action, not mind-reading; make it winnable, and let him win his own way.",
"Armstrong's Appreciation Equation: notice what he already does before you ask for the next thing; give it in his currency; then let him see it landed. That's the fuel.",
"The loop that works: clear invitation, he gives, you receive and show it landed, he wants to give more. Same man, same marriage, you changed how you asked and how you received.",
"Stop deflecting the gift; let the compliment land. You can't ask for intimacy from an empty tank. Receiving is half the skill.",
"Broadcast. In pairs, say one want as a demand first ('You never...'), notice your partner's face; then as an invitation ('I'd so love it if you would...') and show how it would land ('...and it would make me feel...'). Partner says which made them want to give. Swap. Main-room: which were you more used to?",
"Claiming round: 'This week, the one clean want I'm going to actually say is...' Circle answers: You're allowed.",
"Take-home: one clean, specific, appreciation-first request a day; watch the generosity come back; you're teaching the people who love you how to succeed at loving you. Close.",
]

JOBS=[
("content/sessions/slides/Master-Your-Emotions.pptx", EMO),
("content/sessions/slides/Session-1-Coming-Home-to-Yourself.pptx", S1),
("content/sessions/slides/Session-2-Speaking-It-So-Youre-Heard.pptx", S2),
]
for path, notes in JOBS:
    prs=Presentation(path); n=len(prs.slides._sldIdLst)
    assert len(notes)==n, f"{path}: notes {len(notes)} != slides {n}"
    for slide,note in zip(prs.slides, notes):
        slide.notes_slide.notes_text_frame.text=note
    prs.save(path); print("notes ->", path, n)
