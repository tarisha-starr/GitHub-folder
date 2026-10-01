#!/usr/bin/env python3
"""Attach full facilitator-script presenter notes to each slide of the
Spirituality & Sexuality deck, in order."""
from pptx import Presentation

PATH="content/sessions/slides/Spirituality-and-Sexuality.pptx"
prs=Presentation(PATH)

notes=[
# 1 Title
"Welcome everyone. Light the candle. Read the sacred-space agreements: what's shared here stays here; we witness, we don't fix; you can always pass or slow down; there's no such thing as too much in this room; this is a reverent, judgement-free space. Then settle them: eyes closed, feet on the floor, soften jaw, shoulders, belly, three slow breaths down into the pelvis. 'Tonight we do something most of us were never allowed to do: we treat your body, your pleasure and your desire as holy.' One-word go-round: how are you arriving?",
# 2 We begin with self-compassion
"We always begin here. Say it simply: you can't open to anything bigger while you're at war with yourself. So before anything else, we come home to ourselves with kindness.",
# 3 Self-compassion (practice)
"Guide the practice slowly, hands on heart: 'Both hands on your heart. Feel the warmth. Ask yourself, the way you'd ask your dearest friend: How are you, really? What's going on for you? Listen. Then say: I hear you. I'm here. If something tender rises: I see this, and I'm not going anywhere. Nothing to fix. Nothing to perform. Just meeting yourself with kindness.' Sit in it two to three minutes. This is the ground everything stands on.",
# 4 Two steps
"Teach the two steps, in order. Step one is awareness: noticing the moment you close down, contract, go reactive, the old story that says 'not enough.' Step two is self-love. But awareness alone can curdle into self-criticism, 'I can feel it and I still do it, what's wrong with me?' So the key of step two: you don't need to fix yourself. You are not a problem to be solved.",
# 5 You are not a problem to solve
"Let this land. You can feel yourself contracting, closing, bracing, and you don't have to make it go away. Nothing about you needs solving tonight.",
# 6 The practice
"This is the heart of it. The practice is not getting rid of the contraction. The practice is opening to something bigger than you, irrespective of the contraction. You feel yourself close, and you open anyway, not by forcing yourself open, and not by fixing the closing first, but by turning toward something larger while the contraction is still here: love, life, breath, the sacred. Over time the contraction stops running the show, not because you defeated it, but because you stopped making it the centre.",
# 7 Not bypassing
"Name the trap. Spirituality is not bypassing. Bypassing uses 'love and light' to rise above your humanness and pretend the hurt or fear isn't there, which buries it. Real practice does the opposite: it lets the contracted, scared, reactive parts stay exactly as they are, and opens to something bigger while holding them with love. Those parts are usually young and protective. You don't shame them, and you don't fix them. You let them be held.",
# 8 Your body is your Blueprint
"Teach the Blueprint. Your body is always guiding you: it speaks in contraction and opening, yes and no, turn-on and turn-off. That is your Blueprint, your inner guidance system. Most of us were taught to override it and not trust it. Coming home to your body means learning to read and trust those signals again. So when you feel yourself close, that's not a fault, it's information from your guidance system.",
# 9 Safety is the sacred ground
"None of this is sacred without safety. Sacred means you're never pushed, never performing, never abandoning yourself. Your yes is real because your no is allowed. Consent, first with yourself, is the beginning of anything holy. Everything tonight is an invitation: take what's yours, leave the rest.",
# 10 Reflect
"Pose the reflection questions and let them sit in silence for a moment before the breakout: Where do I feel myself contract or close down? What happens when I try to fix it? What could I open to, even while the contraction is still here?",
# 11 Breakout: the part that closes down
"Broadcast to the rooms before opening them. In pairs, Partner A for about 6 minutes: 'I feel myself contract or close down when...' / 'What happens when I try to fix it is...' / 'What I could open to, even while it's there, is...' Partner B simply witnesses, no fixing, for them or for themselves. Then swap. Re-broadcast the swap at halfway. Bring everyone back for a one-sentence go-round: 'What I can open to, even while I'm contracted, is...'",
# 12 What sacred sexuality means
"Teach it warmly. Sacred sexuality means bringing intention, presence and reverence to sex, so it becomes more than a physical act. It's a doorway to connection with something larger than yourself, name it however you like: the divine, life force, oneness, love. The act itself doesn't change. The awareness you bring to it does. Approached with presence, sex becomes communion, not performance or transaction.",
# 13 The same longing
"How they connect: sexuality and spirituality spring from the same root, desire and longing. Sexuality is the longing to know and be known, to merge with and be met by another, body, heart and soul. Spirituality is that same longing pointed at the whole, to dissolve into and belong to something greater. Two directions, one yearning, to end separateness. That's why the old traditions never split them.",
# 14 The body is the doorway
"Tantra never split body and spirit. The Western split, body shameful, spirit holy, is the wound, not the truth. Transcendence, like pleasure, moves through the body, not around it. You reach the sacred by going into the body, not by rising above it. This is exactly why bypassing can't get you there.",
# 15 Presence is the skill (evidence)
"Open the evidence: this isn't wishful thinking. Dr Lori Brotto's clinical studies show mindfulness, simply being present in the body, measurably improves women's desire, arousal and satisfaction, held at six-month follow-up. The spiritual muscle and the sexual muscle are the same muscle: presence. And presence is trainable.",
# 16 Union is literal
"In studies of mutual eye gaze, two people's brain activity, and even their blinking, begin to synchronise, and they feel measurably closer. 'Becoming one' isn't only poetry, there's a biology to it. This is why eye gazing is in the practices.",
# 17 A chemistry of the sacred
"Touch and intimacy release oxytocin, dopamine and endorphins: bonding, trust, dissolved boundaries, the afterglow. That merged, spacious, melted feeling women describe has a real chemistry. Connection literally soothes the nervous system.",
# 18 Vulnerability is the gateway
"Research ties being truly seen to the deepest connection, the very vulnerability the spiritual traditions call surrender. Letting yourself be seen is both the sexual and the spiritual doorway.",
# 19 Your body is a gift
"Bridge it all together. Received as a gift, not a project to fix, and trusted as your guidance, your body isn't the obstacle to the sacred. It's the doorway. A gift is received, not earned. That's the posture of both pleasure and the sacred.",
# 20 Practice: Sacred breath
"Guide slowly, cameras optional. 'Both feet on the floor. One hand on your heart, one hand low on your belly. Breathe down, long and slow, into your womb space, the seat of your life force. Breathe in warmth. Breathe out, and if you feel yourself contract, let it be there, you're not making it leave, you're just opening, breath by breath, to something bigger. Now move your hands slowly, honouring your body. As you touch each part, thank her: what do your hands do for you, your hips, your belly? Send each part gratitude and love. Silently: I'm here. I love you. Thank you.' Pleasure opens in a body that feels adored.",
# 21 Breakout: From shame to sacred
"Pose first: what pleasure did I stop letting myself feel, and what would it mean to give it back? Then broadcast. In pairs, Partner A for about 6 minutes: 'A pleasure I stopped letting myself feel is...' / 'One small, sacred pleasure I could give myself this week is...' / 'The part of my body I'm ready to make peace with is...' Partner B witnesses with warmth. Swap. Main-room: name the one sacred pleasure you're claiming this week.",
# 22 This week (take-home)
"Send them off. Each day this week, ten minutes: light a candle, hand on heart and hand on belly, breathe into the body and say 'I'm here. I love you. Thank you.' Then do one small thing that treats the body as sacred, a slow bath, silk on the skin, moving to music, dressing for no one but yourself. Closing round, one line each: 'The way I'll honour my body this week is...' Circle answers: 'She is sacred.' Blow out the candle together.",
# 23 Session 2 title
"Session 2. Soften them in as before, then recap the ground: last time we made the body sacred again; tonight we bring that reverence to connection; safety first, always, sacred union begins with your own yes. Go-round: one way I honoured my body this week. Pose the reflection questions for later: When do I feel truly safe and connected? What do I long for that I've never said out loud? What helps me stay present instead of performing?",
# 24 Slow is sacred
"Slow is the secret. Most lovemaking rushes to stimulation and a goal, but you can't build deep connection at full speed. Sacred union isn't performance or arriving anywhere, it's presence. Make love the way you'd meditate or pray: fully here, fully awake, in no hurry at all.",
# 25 Presence, not performance
"When you slow right down, something opens: space for feeling, space for vulnerability. And vulnerability is what lets two people actually meet, instead of two bodies going through the motions. Slowness turns sex into communion. Feeling is what makes it sacred.",
# 26 Reflect
"Pose and let land before the practice and breakout: When do I feel truly safe and connected with another, and what creates it? What do I long for in intimacy that I've never said out loud? What helps me stay present instead of performing or disappearing?",
# 27 Practice: Eye gaze & breath
"Guide it. On Zoom they do it solo as rehearsal: hand on heart, or a soft gaze at their own eyes. 'Sit tall. Soften your gaze. Breathe slowly, in and out. Imagine sitting knee to knee with your beloved. You breathe together. You simply witness and are witnessed. No talking, no fixing, no performing. Notice the urge to look away or laugh, to break it. Stay anyway. Underneath the awkwardness there's a person who can be truly seen. Tonight, practise letting yourself be seen, starting with yourself.'",
# 28 Breakout: Safe and connected
"Broadcast. In pairs, Partner A about 6 minutes: 'I feel most safe and connected with someone when...' / 'What pulls me out of presence in intimacy is...' / 'What helps me come back is...' Partner B witnesses. Swap. Main-room: one thing that helps you feel safe and connected.",
# 29 Create sacred space
"Lovemaking can be a ritual, and it starts before you ever touch. Create a special space: candles, soft music, fabrics that feel beautiful, a door that closes. You're telling your nervous system, and your beloved: this is sacred time.",
# 30 The dance of love
"The dance of love, and it's slow. Breathe together first. Eye gaze. See each other's beauty, power and vulnerability. Then touch, slowly, feather the whole body, fingertips sending love, no destination. Take turns giving and receiving. Frame it as presence, not technique: the point isn't what you do, it's how present you are while you do it.",
# 31 Breakout: Desire & receiving
"Pose first: what do I desire that I've been too afraid to ask for, and what would it feel like to simply receive? Broadcast. In pairs, Partner A about 6 minutes, practising two things out loud: speak a desire cleanly and warmly, 'Something I long for in intimacy is...' (a desire, not a complaint); then practise receiving, let Partner B offer a genuine appreciation and simply say 'thank you,' let it land, notice the urge to deflect. Partner B witnesses and offers the appreciation. Swap. Main-room: what was harder, speaking your desire or receiving?",
# 32 This week (couple's ritual)
"Take-home, solo or with a chosen partner: create one pocket of sacred time. Light a candle. Sit knee to knee, or hand on heart if solo. Breathe together a few minutes. Eye gaze for three. Then slow, reverent touch or self-touch, no goal at all, just presence and love. Notice how different it feels when you're fully here.",
# 33 You are sacred (close)
"Close the circle. Claiming round, one line each: 'The reverence I'm bringing to intimacy is...' The circle answers together, warm: 'Sacred.' Remind them: your body, your pleasure, your desire, all holy. Blow out the candle together.",
]

n=len(prs.slides._sldIdLst)
assert len(notes)==n, f"notes {len(notes)} != slides {n}"
for slide,note in zip(prs.slides, notes):
    slide.notes_slide.notes_text_frame.text=note
prs.save(PATH)
print("added notes to", n, "slides")
