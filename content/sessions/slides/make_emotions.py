#!/usr/bin/env python3
"""Master Your Emotions deck (safety-first). Self-contained so it does not
regenerate the other decks. Same brand style: full-bleed palette backgrounds,
Marcellus headings, Lora body, upright (no italics), one line or a few bullets."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
import os

CREAM=RGBColor(0xF5,0xEF,0xE3); IVORY=RGBColor(0xFB,0xF7,0xEE)
BURGUNDY=RGBColor(0x6E,0x1A,0x2E); NAVY=RGBColor(0x1F,0x2A,0x44)
RUST=RGBColor(0x9E,0x4A,0x2A); COPPER=RGBColor(0xC7,0x5D,0x3D)
GOLD=RGBColor(0xC2,0xA4,0x6D); NEARBLK=RGBColor(0x15,0x11,0x0D)

HEAD_FONT="Marcellus"; BODY_FONT="Lora"

SCHEME={"cream":(CREAM,NEARBLK,COPPER),"ivory":(IVORY,NEARBLK,COPPER),
        "burgundy":(BURGUNDY,CREAM,GOLD),"navy":(NAVY,CREAM,GOLD),
        "rust":(RUST,CREAM,GOLD),"copper":(COPPER,CREAM,CREAM),
        "gold":(GOLD,NEARBLK,BURGUNDY)}


def new_deck():
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    return prs


def add_slide(prs,bg,title,body=None,eyebrow=None,big=False,footer=True):
    bg_col,txt_col,accent=SCHEME[bg]
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=bg_col
    left=Inches(0.9); width=Inches(11.5); top=Inches(1.1)
    if eyebrow:
        eb=s.shapes.add_textbox(left,Inches(0.6),width,Inches(0.5))
        r=eb.text_frame.paragraphs[0].add_run(); r.text=eyebrow.upper()
        r.font.name=BODY_FONT; r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=accent
    tb=s.shapes.add_textbox(left,top,width,Inches(3.2))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    r=tf.paragraphs[0].add_run(); r.text=title
    r.font.name=HEAD_FONT; r.font.size=Pt(54 if big else 40); r.font.color.rgb=txt_col
    line=s.shapes.add_shape(1,left,top+Inches(3.25 if big else 2.0),Inches(1.6),Pt(3))
    line.fill.solid(); line.fill.fore_color.rgb=accent; line.line.fill.background(); line.shadow.inherit=False
    if body:
        by=s.shapes.add_textbox(left,top+Inches(3.6 if big else 2.35),width,Inches(3.2))
        bf=by.text_frame; bf.word_wrap=True
        items=body if isinstance(body,list) else [body]
        for i,item in enumerate(items):
            para=bf.paragraphs[0] if i==0 else bf.add_paragraph()
            bullet=isinstance(body,list) and len(body)>1
            run=para.add_run(); run.text=("•  "+item) if bullet else item
            run.font.name=BODY_FONT; run.font.size=Pt(23 if bullet else 27)
            run.font.italic=False; run.font.color.rgb=txt_col
            para.space_after=Pt(12)
    if footer:
        ft=s.shapes.add_textbox(left,Inches(6.85),width,Inches(0.4))
        fr=ft.text_frame.paragraphs[0].add_run(); fr.text="SEXUALEMPOWERMENTFORWOMEN.COM"
        fr.font.name=BODY_FONT; fr.font.size=Pt(11); fr.font.color.rgb=accent
    return s


d=new_deck()

add_slide(d,"burgundy","Master Your Emotions",
          eyebrow="Radiant Women's Circle",
          body="Feeling into them, without numbing and without drowning.",big=True)

add_slide(d,"cream","Lack of safety is the problem.",
          body="Not the emotions. Emotions aren't dangerous. Being flooded is.",big=True)

add_slide(d,"rust","Start with safety.",
          body="We only expand, feel more, open more, when we're not overwhelmed and not sliding into our old patterns.")

add_slide(d,"navy","What safety is",
          body=["I'm connected to myself.",
                "I feel what's good for me.",
                "I test it out.",
                "I read how the other responds.",
                "Then I step forward, or step away."])

add_slide(d,"gold","Then, no overwhelm.",
          body="Because I never abandon myself to cope. I stay with me the whole way.")

add_slide(d,"burgundy","The loop",
          body="Stress makes emotions. Emotions make stress. So we reach for something to fill the hole.")

add_slide(d,"cream","The hole can't be filled from outside",
          body="There's no signal that says your emotional need is met. The fix backfires, and the feeling still waits to be heard.")

add_slide(d,"rust","Reflection",
          eyebrow="One minute",
          body="What do you reach for when a feeling gets too big?")

add_slide(d,"navy","The reframe",
          eyebrow="Dr Linda Bacon",
          body="If you numb, you don't have a problem with food or wine or your phone. You have a problem with taking care of yourself.")

add_slide(d,"gold","Self-compassion",
          eyebrow="Hands on heart",
          body=["How do you feel? What's going on for you?",
                "I hear you feel...",
                "I see this is difficult. I am here."])

add_slide(d,"burgundy","Riding the wave",
          body="Grief comes in waves. Feel a little, come back to safety, feel a little more. You don't have to feel all of it at once.")

add_slide(d,"cream","Feel without being flooded",
          body=["Feel a little, then come back.",
                "Keep one foot in the present.",
                "Let the tears come. They're a release.",
                "Then soothe. You did something brave."])

add_slide(d,"navy","The body still braces",
          body="The mind makes peace in words. The body lets go in a different language: breath, tears, movement, touch.")

add_slide(d,"rust","EFT · tapping",
          body="An easy, direct way to settle a big feeling and come back to safety in your body.")

add_slide(d,"gold","Why it works",
          body="Tapping calms the amygdala's alarm, so you shift into your logical brain and choose your response.")

add_slide(d,"cream","Wash the laundry",
          body="We don't hide the feeling in the wardrobe. We take it out, wash it, and clear it. Then reframes come naturally.")

add_slide(d,"navy","How to tap",
          body=["Name it, be specific. Rate it 1 to 10.",
                "Setup: even though I feel this, I love and accept myself.",
                "Tap the points, say it out loud.",
                "Breathe, re-rate, repeat until it drops."])

add_slide(d,"rust","The points",
          body=["Top of head · eyebrow · side of eye · under eye",
                "Under nose · chin · collarbone",
                "Under arm · wrist",
                "Sip water as you go."])

add_slide(d,"burgundy","This week",
          eyebrow="Take home",
          body=["Safety first: am I connected to myself?",
                "Hands on heart: how do you feel?",
                "Tap when it's big. Feel a little, come back."])

add_slide(d,"gold","You don't control your emotions.",
          body="You master them by feeling safe enough to finally listen.",big=True)

out="content/sessions/slides/Master-Your-Emotions.pptx"
os.makedirs("content/sessions/slides",exist_ok=True)
d.save(out)
print("saved",out,len(d.slides._sldIdLst),"slides")
