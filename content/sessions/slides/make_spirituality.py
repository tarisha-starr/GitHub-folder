#!/usr/bin/env python3
"""Spirituality & Sexuality - one combined PowerPoint (both Zoom sessions).
Self-contained. SEFW palette only, Marcellus + Lora, upright text."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
import os

NAVY =RGBColor(0x1D,0x28,0x3C); PLUM =RGBColor(0x74,0x23,0x4F)
BLUSH=RGBColor(0xE4,0xBB,0xC2); PINK =RGBColor(0xFC,0xE8,0xEA)
SAGE =RGBColor(0x7C,0xA1,0xA5); TEAL =RGBColor(0x22,0x46,0x52)
GOLD =RGBColor(0xC9,0xA8,0x6D)
HEAD_FONT="Marcellus"; BODY_FONT="Lora"
SCHEME={"pink":(PINK,NAVY,PLUM),"blush":(BLUSH,NAVY,PLUM),"plum":(PLUM,PINK,GOLD),
        "navy":(NAVY,PINK,GOLD),"teal":(TEAL,PINK,GOLD),"sage":(SAGE,NAVY,PLUM),
        "gold":(GOLD,NAVY,PLUM)}

def new_deck():
    p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5); return p

def add(prs,bg,title,body=None,eyebrow=None,big=False,footer=True):
    bg_col,txt_col,accent=SCHEME[bg]
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=bg_col
    left=Inches(0.9); width=Inches(11.5); top=Inches(1.1)
    if eyebrow:
        eb=s.shapes.add_textbox(left,Inches(0.6),width,Inches(0.5))
        r=eb.text_frame.paragraphs[0].add_run(); r.text=eyebrow.upper()
        r.font.name=BODY_FONT; r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=accent
    tb=s.shapes.add_textbox(left,top,width,Inches(3.2))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    r=tf.paragraphs[0].add_run(); r.text=title
    r.font.name=HEAD_FONT; r.font.size=Pt(54 if big else 40); r.font.color.rgb=txt_col
    ln=s.shapes.add_shape(1,left,top+Inches(3.25 if big else 2.0),Inches(1.6),Pt(3))
    ln.fill.solid(); ln.fill.fore_color.rgb=accent; ln.line.fill.background(); ln.shadow.inherit=False
    if body:
        by=s.shapes.add_textbox(left,top+Inches(3.6 if big else 2.35),width,Inches(3.2))
        bf=by.text_frame; bf.word_wrap=True
        items=body if isinstance(body,list) else [body]
        for i,item in enumerate(items):
            para=bf.paragraphs[0] if i==0 else bf.add_paragraph()
            bullet=isinstance(body,list) and len(body)>1
            run=para.add_run(); run.text=("•  "+item) if bullet else item
            run.font.name=BODY_FONT; run.font.size=Pt(23 if bullet else 27)
            run.font.italic=False; run.font.color.rgb=txt_col
            para.space_after=Pt(12)
    if footer:
        ft=s.shapes.add_textbox(left,Inches(6.85),width,Inches(0.4))
        fr=ft.text_frame.paragraphs[0].add_run(); fr.text="SEXUALEMPOWERMENTFORWOMEN.COM"
        fr.font.name=BODY_FONT; fr.font.size=Pt(11); fr.font.color.rgb=accent
    return s

d=new_deck()

# ---- Session 1: The Sacred Woman ----
add(d,"navy","Spirituality & Sexuality",eyebrow="Radiant Women's Circle  ·  Session 1",
    body="The Sacred Woman: coming home to your body and desire.",big=True)
add(d,"pink","Sacred, not shameful",
    body="Your sexuality isn't the opposite of your spirituality. It's one of its deepest doorways.",big=True)
add(d,"plum","The body is a temple",
    body="Pleasure is a prayer. Your sexual energy is your life force, the same energy that creates life.")
add(d,"navy","Safety is the sacred ground",
    body="Nothing is pushed. You go at your own pace. Your yes is real because your no is allowed.")
add(d,"blush","Reflect",eyebrow="Take a moment",
    body=["What was I taught about my body and sexuality?","Where do I still carry shame? Whose voice is it?",
          "What would it feel like to treat my body as sacred?"])
add(d,"sage","Practice · Sacred breath",
    body="Hand on heart, hand on belly. Breathe down into your womb space. I'm here. I love you. Thank you.")
add(d,"teal","Breakout · The story I was told",eyebrow="In pairs",
    body=["The story I was told about my body was...","The shame I carry is... in the voice of...",
          "What I'd love to believe instead is...","Partner: witness. Then swap."])
add(d,"gold","Your body is a temple",
    body="Pleasure can't open in a body that's bracing for criticism. It opens in a body that feels adored.")
add(d,"blush","Breakout · From shame to sacred",eyebrow="In pairs",
    body=["A pleasure I stopped letting myself feel is...","One sacred pleasure I'll give myself is...",
          "The part of my body I'm ready to make peace with is...","Then swap."])
add(d,"navy","This week",eyebrow="Take home",
    body=["Candle. Hand on heart, hand on belly.","I'm here. I love you. Thank you.",
          "One small thing that treats your body as sacred."])

# ---- Session 2: The Sacred Union ----
add(d,"navy","Spirituality & Sexuality",eyebrow="Radiant Women's Circle  ·  Session 2",
    body="The Sacred Union: deepening connection during lovemaking.",big=True)
add(d,"plum","Slow is sacred",
    body="You can't build deep connection at full speed. Make love the way you'd meditate: fully here, in no hurry.",big=True)
add(d,"pink","Presence, not performance",
    body="Sacred union isn't about arriving anywhere. Slowness makes space for feeling, and feeling makes it sacred.")
add(d,"blush","Reflect",eyebrow="Take a moment",
    body=["When do I feel truly safe and connected?","What do I long for that I've never said out loud?",
          "What helps me stay present instead of performing?"])
add(d,"sage","Practice · Eye gaze & breath",
    body="Soften your gaze. Breathe together. Simply witness and be witnessed. Let yourself be seen.")
add(d,"teal","Breakout · Safe and connected",eyebrow="In pairs",
    body=["I feel most safe and connected when...","What pulls me out of presence is...",
          "What helps me come back is...","Then swap."])
add(d,"navy","Create sacred space",
    body="Candles. Soft music. Beautiful fabrics. A door that closes. You're telling your body: this is sacred time.")
add(d,"gold","The dance of love",
    body="Breathe together. Eye gaze. Then touch, slowly, with no destination. Take turns giving and receiving.")
add(d,"blush","Breakout · Desire & receiving",eyebrow="In pairs",
    body=["Speak a desire, cleanly and warmly: I long for...","Receive an appreciation. Just say thank you. Let it land.",
          "Notice the urge to deflect.","Then swap."])
add(d,"plum","This week",eyebrow="Take home",
    body=["One pocket of sacred time. Light a candle.","Breathe together. Eye gaze for three minutes.",
          "Slow, reverent touch, with no goal. Just presence."])
add(d,"navy","You are sacred.",
    body="Your body. Your pleasure. Your desire. All holy.",big=True)

os.makedirs("content/sessions/slides",exist_ok=True)
out="content/sessions/slides/Spirituality-and-Sexuality.pptx"
d.save(out)
print("saved",out,len(d.slides._sldIdLst),"slides")
